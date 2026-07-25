from __future__ import annotations

import hashlib
import json
import re
import statistics
import sys
import warnings
from collections import Counter, defaultdict
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
DATA_ROOT = ROOT / "06.分级系统数据源"
RUN_ROOT = Path(__file__).resolve().parent
NORMALIZED = RUN_ROOT / "normalized_inputs"
SELLER_WORKBOOK = DATA_ROOT / "插件数据源_女士休闲连衣裙_卖家精灵全量信息_20260710.xlsx"
FRONTEND_ROOT = DATA_ROOT / "前端数据源最终报表_按站点"
KALODATA_ENVELOPE = (
    DATA_ROOT / "Kalodata连衣裙系统数据源" / "Kalodata连衣裙_20260710.source_envelope.json"
)
COLLECTED_AT = datetime.now(timezone.utc).isoformat(timespec="seconds")


def present(value: Any) -> bool:
    return value not in (None, "", [], {})


def serializable(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, tuple):
        return list(value)
    return value


def rel(path: Path) -> str:
    return str(path.resolve().relative_to(ROOT.resolve()))


def slug(text: str) -> str:
    value = re.sub(r"[^a-z0-9]+", "_", text.lower()).strip("_")
    return value or hashlib.sha1(text.encode()).hexdigest()[:12]


def number(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    match = re.search(r"-?\d+(?:,\d{3})*(?:\.\d+)?", str(value))
    return float(match.group().replace(",", "")) if match else None


def percent_points(value: Any) -> float | None:
    parsed = number(value)
    if parsed is None:
        return None
    if isinstance(value, str) and "%" in value:
        return parsed / 100
    return parsed / 100 if abs(parsed) > 2 else parsed


def money(value: Any, currency: Any = None) -> dict[str, Any] | None:
    amount = number(value)
    if amount is None:
        return None
    code = str(currency or "").strip().upper()
    text = str(value)
    if not code:
        code = "USD" if "$" in text else "CNY" if "¥" in text or "￥" in text else ""
    return {"amount": amount, "currency": code} if code else None


def sheet_rows(ws):
    iterator = ws.iter_rows(values_only=True)
    headers = [str(value).strip() if value is not None else "" for value in next(iterator)]
    for row_number, row in enumerate(iterator, 2):
        yield row_number, {headers[index]: serializable(value) for index, value in enumerate(row)}


def evidence(source_id: str, source_type: str, artifact: Path, row: int,
             field: str, column: str, confidence: str = "high", note: str = ""):
    digest = hashlib.sha1(f"{source_id}|{row}|{field}|{column}".encode()).hexdigest()[:16]
    return {
        "evidence_id": f"ev_{digest}",
        "source_id": source_id,
        "source_type": source_type,
        "tool_or_adapter": "full_selection_input_normalizer/1.0",
        "field_path": field,
        "artifact_path": rel(artifact),
        "record_locator": f"sheet=商品详情;row={row};column={column}",
        "observed_at": COLLECTED_AT,
        "confidence": confidence,
        "note": note,
    }


def add_field(record: dict[str, Any], source_id: str, source_type: str,
              artifact: Path, row: int, field: str, value: Any, column: str,
              confidence: str = "high", note: str = ""):
    if not present(value):
        return
    record["fields"][field] = serializable(value)
    record["evidence"][field] = evidence(
        source_id, source_type, artifact, row, field, column, confidence, note)


def series_growth(values: list[float]) -> float | None:
    values = [float(value) for value in values if value is not None]
    if len(values) < 4:
        return None
    head = statistics.fmean(values[: min(3, len(values))])
    tail = statistics.fmean(values[-min(3, len(values)):])
    return tail / head - 1 if head > 0 else None


def aggregate_sheet(ws, key_name: str, value_names: list[str], skip_current: str | None = None):
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    rows = 0
    for _, record in sheet_rows(ws):
        rows += 1
        key = str(record.get(key_name) or "").strip()
        if not key:
            continue
        if skip_current and record.get(skip_current) is True:
            continue
        grouped[key].append({name: record.get(name) for name in value_names})
    return grouped, rows


def seller_trends(workbook):
    sales, sales_rows = aggregate_sheet(
        workbook["销量趋势"], "ASIN", ["月份", "销量", "销售额($)"], "是否当前月")
    bsr, bsr_rows = aggregate_sheet(workbook["BSR趋势"], "ASIN", ["日期", "BSR排名"])
    prices, price_rows = aggregate_sheet(workbook["价格趋势"], "ASIN", ["日期", "价格($)"])
    reviews, review_rows = aggregate_sheet(workbook["评论趋势"], "ASIN", ["日期", "评分数", "星级"])
    summaries = {}
    for asin in set(sales) | set(bsr) | set(prices) | set(reviews):
        unit_values = [number(row.get("销量")) for row in sales.get(asin, [])]
        bsr_values = [number(row.get("BSR排名")) for row in bsr.get(asin, [])]
        price_values = [number(row.get("价格($)")) for row in prices.get(asin, [])]
        review_values = [number(row.get("评分数")) for row in reviews.get(asin, [])]
        unit_values = [value for value in unit_values if value is not None]
        bsr_values = [value for value in bsr_values if value and value > 0]
        price_values = [value for value in price_values if value is not None]
        review_values = [value for value in review_values if value is not None]
        summary = {
            "sales_points": len(unit_values),
            "sales_growth_full_history": series_growth(unit_values),
            "bsr_points": len(bsr_values),
            "bsr_best": min(bsr_values) if bsr_values else None,
            "bsr_latest": bsr_values[-1] if bsr_values else None,
            "price_points": len(price_values),
            "price_min": min(price_values) if price_values else None,
            "price_max": max(price_values) if price_values else None,
            "price_cv": (
                statistics.pstdev(price_values) / statistics.fmean(price_values)
                if len(price_values) > 1 and statistics.fmean(price_values) else None),
            "review_points": len(review_values),
            "review_gain": review_values[-1] - review_values[0] if len(review_values) > 1 else None,
        }
        tags = []
        if summary["sales_growth_full_history"] is not None:
            if summary["sales_growth_full_history"] >= 0.30:
                tags.append("sales_trend_rising")
            elif summary["sales_growth_full_history"] >= -0.10:
                tags.append("sales_trend_stable")
        if len(bsr_values) >= 8:
            first = statistics.median(bsr_values[:4])
            last = statistics.median(bsr_values[-4:])
            if last < first * 0.7:
                tags.append("bsr_improving")
        if summary["price_cv"] is not None and summary["price_cv"] <= 0.05:
            tags.append("price_stable")
        summary["derived_tags"] = tags
        summaries[asin] = summary
    return summaries, {
        "销量趋势": sales_rows,
        "BSR趋势": bsr_rows,
        "价格趋势": price_rows,
        "评论趋势": review_rows,
    }


def seller_raw_fields(workbook):
    keep = {
        "allImages", "video", "variation_list_all", "features", "product_attributes",
        "five_bullet_points", "location", "fulfillment", "dimensions", "weight",
        "pkg_dimensions", "pkg_weight", "pkg_dimension_type", "delivery_price",
        "coupon", "prime_exclusive_price", "aplus_content", "ebc", "qa",
    }
    selected: dict[str, dict[str, Any]] = defaultdict(dict)
    field_names: dict[str, Counter] = defaultdict(Counter)
    rows = 0
    for _, row in sheet_rows(workbook["全字段长表"]):
        rows += 1
        asin = str(row.get("ASIN") or "").strip()
        name = str(row.get("原字段名") or "").strip()
        if not asin or not name:
            continue
        field_names[asin][name] += 1
        if name in keep and name not in selected[asin] and present(row.get("原始值")):
            selected[asin][name] = row.get("原始值")
    return selected, field_names, rows


def seller_envelope():
    warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")
    workbook = load_workbook(SELLER_WORKBOOK, read_only=True, data_only=True)
    trends, trend_rows = seller_trends(workbook)
    raw_selected, raw_names, raw_rows = seller_raw_fields(workbook)
    source_id = "seller_sprite_amazon"
    records = []
    for row_number, row in sheet_rows(workbook["产品总表"]):
        asin = str(row.get("ASIN") or "").strip()
        if not asin:
            continue
        record = {
            "record_kind": "product",
            "platform": "amazon",
            "site": "amazon",
            "product_id": asin,
            "id_truncated": False,
            "category_label": "women_casual_dresses",
            "seasonal": False,
            "quality_status": "VALID",
            "fields": {},
            "evidence": {},
        }
        def add(field, value, column, confidence="high", note=""):
            add_field(record, source_id, "plugin", SELLER_WORKBOOK, row_number,
                      field, value, column, confidence, note)

        category_path = [part.strip() for part in str(row.get("类目路径") or "").split(":") if part.strip()]
        add("title", row.get("英文标题"), "英文标题")
        add("title_cn", row.get("中文标题"), "中文标题")
        add("brand", row.get("品牌"), "品牌")
        add("category_name", category_path[-1] if category_path else "Women Casual Dresses", "类目路径")
        add("category_path", category_path, "类目路径")
        add("canonical_url", row.get("Amazon链接"), "Amazon链接")
        add("image_url", row.get("图片链接"), "图片链接")
        add("unit_price", money(row.get("价格($)"), "USD"), "价格($)")
        add("rating", number(row.get("星级")), "星级")
        add("review_count", number(row.get("评分数")), "评分数")
        add("sales_30d_units", number(row.get("近30天销量(父体)")), "近30天销量(父体)",
            "medium", "SellerSprite 估算；父体近30天销量")
        add("revenue_30d", money(row.get("近30天销售额($)"), "USD"), "近30天销售额($)",
            "medium", "SellerSprite 估算")
        growth = percent_points(row.get("销量增长率"))
        add("sales_growth_rate", growth, "销量增长率", "medium", "源表百分数口径归一为小数")
        add("bestseller_rank_text", row.get("Amazon类目排名"), "Amazon类目排名")
        add("seller_name", row.get("卖家名称"), "卖家名称")
        add("image_count", number(row.get("主图数量")), "主图数量")
        add("review_tags_raw", row.get("评论亮点"), "评论亮点", "medium")
        add("raw_evidence_path", rel(SELLER_WORKBOOK), "采集说明")
        design_signals = []
        for kind, column in (
            ("bullet_points", "英文五点"), ("material", "面料"),
            ("attributes", "产品属性JSON"), ("review_summary", "AI评论摘要"),
        ):
            if present(row.get(column)):
                design_signals.append({"kind": kind, "text": str(row[column])[:4000]})
        add("design_signals", design_signals, "英文五点/面料/产品属性JSON", "high")
        summary = trends.get(asin, {})
        add("trend_tags", summary.get("derived_tags"), "销量趋势/BSR趋势/价格趋势", "medium",
            "全量趋势行聚合得出")
        for field, value, column in (
            ("description_text", row.get("英文五点"), "英文五点"),
            ("description_text_cn", row.get("中文五点"), "中文五点"),
            ("material_text", row.get("面料"), "面料"),
            ("product_attributes_json", row.get("产品属性JSON"), "产品属性JSON"),
            ("parent_asin", row.get("父体ASIN"), "父体ASIN"),
            ("variation_count", number(row.get("变体数")), "变体数"),
            ("first_available_date", row.get("首次上架日期"), "首次上架日期"),
            ("package_dimensions", row.get("包装尺寸"), "包装尺寸"),
            ("package_weight", row.get("包装重量"), "包装重量"),
            ("item_weight", row.get("商品重量"), "商品重量"),
            ("item_dimensions", row.get("商品尺寸"), "商品尺寸"),
            ("package_type", row.get("包装类型"), "包装类型"),
            ("seller_location", row.get("卖家所属地"), "卖家所属地"),
            ("seller_fulfillment", row.get("卖家类型"), "卖家类型"),
            ("has_aplus", row.get("是否有A+"), "是否有A+"),
            ("has_video", row.get("是否有视频"), "是否有视频"),
            ("seller_sprite_trend_summary", summary, "四个趋势工作表"),
            ("seller_sprite_raw_field_names", sorted(raw_names.get(asin, {})), "全字段长表"),
            ("seller_sprite_selected_raw", raw_selected.get(asin), "全字段长表"),
        ):
            add(field, value, column, "medium" if "trend" in field or "raw" in field else "high")
        records.append(record)
    workbook_sheet_rows = {
        name: max((workbook[name].max_row or 1) - 1, 0) for name in workbook.sheetnames
    }
    workbook.close()
    envelope = {
        "schema_version": "source.envelope.v1",
        "envelope_id": "envelope_seller_sprite_amazon_20260710",
        "source_id": source_id,
        "source_type": "plugin",
        "market": "US",
        "collection_mode": "full_then_analyze",
        "layer_hint": "L1",
        "collected_at": COLLECTED_AT,
        "collector_version": "full_selection_input_normalizer/1.0",
        "quality_status": "VALID",
        "records": records,
        "errors": [],
        "warnings": ["销量与销售额为 SellerSprite 估算值；趋势工作表按全部有效行聚合。"],
        "artifact_paths": {"workbook": rel(SELLER_WORKBOOK)},
    }
    audit = {
        "source_id": source_id,
        "source_type": "plugin",
        "artifact_path": rel(SELLER_WORKBOOK),
        "candidate_records": len(records),
        "sheet_rows_scanned": workbook_sheet_rows,
        "trend_rows_consumed": trend_rows,
        "full_field_rows_consumed": raw_rows,
        "usage": "全部产品行进入评分；全部趋势行聚合；全字段长表逐行扫描并保留字段目录与关键原值。",
    }
    return envelope, audit


def frontend_envelope(path: Path):
    site = slug(path.stem)
    source_id = f"frontend_{site}"
    workbook = load_workbook(path, read_only=True, data_only=True)
    records = []
    partial = 0
    skipped = 0
    for row_number, row in sheet_rows(workbook["商品详情"]):
        product_id = str(row.get("SKU/商品ID") or "").strip()
        url = str(row.get("商品URL") or "").strip()
        if not product_id and url:
            product_id = "url_" + hashlib.sha1(url.encode()).hexdigest()[:16]
        if not product_id and not url:
            skipped += 1
            continue
        status_text = str(row.get("详情采集状态") or "").lower()
        quality = "PARTIAL_SOURCE_MISSING" if any(
            token in status_text for token in ("fail", "blocked", "challenge", "error")) else "VALID"
        partial += quality != "VALID"
        record = {
            "record_kind": "product",
            "platform": site,
            "site": site,
            "product_id": product_id,
            "id_truncated": False,
            "category_label": "women_dresses_multisite",
            "seasonal": False,
            "quality_status": quality,
            "fields": {},
            "evidence": {},
        }
        def add(field, value, column, confidence="high", note=""):
            add_field(record, source_id, "frontend", path, row_number,
                      field, value, column, confidence, note)

        price = money(row.get("价格"), row.get("币种"))
        add("title", row.get("商品名称"), "商品名称")
        add("brand", row.get("品牌"), "品牌")
        add("category_name", "Women Dresses", "连衣裙类目URL", "medium", "由交付目录品类范围确定")
        add("canonical_url", url, "商品URL")
        add("image_url", row.get("主图URL"), "主图URL")
        add("unit_price", price, "价格/币种")
        add("rating", number(row.get("评分")), "评分")
        add("review_count", number(row.get("评价数")), "评价数")
        add("raw_evidence_path", rel(path), "详情采集状态")
        design_signals = []
        for kind, column in (
            ("description", "商品描述"), ("material", "材质/面料"),
            ("care", "洗护说明"), ("color", "颜色"), ("size", "尺码"),
            ("details", "商品详细信息"),
        ):
            if present(row.get(column)):
                design_signals.append({"kind": kind, "text": str(row[column])[:4000]})
        add("design_signals", design_signals, "商品描述/材质/洗护/颜色/尺码/商品详细信息")
        for field, column in (
            ("description_text", "商品描述"), ("material_text", "材质/面料"),
            ("care_text", "洗护说明"), ("color_text", "颜色"),
            ("size_text", "尺码"), ("product_details_text", "商品详细信息"),
            ("category_source", "连衣裙类目URL"), ("detail_status", "详情采集状态"),
        ):
            add(field, row.get(column), column)
        records.append(record)
    sheet_counts = {name: max((workbook[name].max_row or 1) - 1, 0) for name in workbook.sheetnames}
    workbook.close()
    quality = "PARTIAL_SOURCE_MISSING" if partial or not records else "VALID"
    warnings_list = []
    if partial:
        warnings_list.append(f"{partial} 行详情采集状态为失败或受阻，保留列表事实并降级。")
    if not records:
        warnings_list.append("该站点本轮无有效商品行，仍保留数据源审计记录。")
    envelope = {
        "schema_version": "source.envelope.v1",
        "envelope_id": f"envelope_{source_id}_20260710",
        "source_id": source_id,
        "source_type": "frontend",
        "market": "US",
        "collection_mode": "full_then_analyze",
        "layer_hint": "L1",
        "collected_at": COLLECTED_AT,
        "collector_version": "full_selection_input_normalizer/1.0",
        "quality_status": quality,
        "records": records,
        "errors": [],
        "warnings": warnings_list,
        "artifact_paths": {"workbook": rel(path)},
    }
    audit = {
        "source_id": source_id,
        "source_type": "frontend",
        "artifact_path": rel(path),
        "candidate_records": len(records),
        "partial_records": partial,
        "skipped_records": skipped,
        "sheet_rows_scanned": sheet_counts,
        "usage": "商品详情全部有效行进入评分或详情增强；失败行保留可见事实并降级，不补造缺失值。",
    }
    return envelope, audit


def kalodata_audit():
    envelope = json.loads(KALODATA_ENVELOPE.read_text(encoding="utf-8"))
    package = KALODATA_ENVELOPE.parent
    jsonl_counts = {}
    for path in sorted(package.rglob("*.jsonl")):
        with path.open(encoding="utf-8") as handle:
            jsonl_counts[rel(path)] = sum(1 for line in handle if line.strip())
    manifest = json.loads((package / "metadata/delivery_manifest.json").read_text(encoding="utf-8"))
    return {
        "source_id": envelope["source_id"],
        "source_type": envelope["source_type"],
        "artifact_path": rel(KALODATA_ENVELOPE),
        "candidate_records": len(envelope.get("records") or []),
        "package_files_verified": manifest.get("file_count_excluding_this_manifest"),
        "jsonl_rows_available": jsonl_counts,
        "usage": "93 个标准商品记录全量进入评分；原始响应、处理表和账本通过交付清单与哈希保留审计。",
        "warnings": envelope.get("warnings") or [],
    }


def main():
    NORMALIZED.mkdir(parents=True, exist_ok=True)
    for old in NORMALIZED.glob("*.json"):
        old.unlink()
    audits = []
    seller, audit = seller_envelope()
    (NORMALIZED / "seller_sprite_amazon.json").write_text(
        json.dumps(seller, ensure_ascii=False, indent=2, default=serializable), encoding="utf-8")
    audits.append(audit)
    for workbook in sorted(FRONTEND_ROOT.glob("*.xlsx")):
        if workbook.name.startswith("00_"):
            continue
        envelope, audit = frontend_envelope(workbook)
        (NORMALIZED / f"{envelope['source_id']}.json").write_text(
            json.dumps(envelope, ensure_ascii=False, indent=2, default=serializable), encoding="utf-8")
        audits.append(audit)
    audits.append(kalodata_audit())
    summary = {
        "generated_at": COLLECTED_AT,
        "input_root": rel(DATA_ROOT),
        "sources": audits,
        "source_count": len(audits),
        "candidate_records_before_cross_source_merge": sum(
            item.get("candidate_records", 0) for item in audits),
        "frontend_workbooks": sum(item["source_type"] == "frontend" for item in audits),
        "plugin_sources": sum(item["source_type"] == "plugin" for item in audits),
        "full_use_policy": (
            "候选级数据全量进入分级；时间序列全量聚合；未知字段保留在标准信封和原始文件审计中；"
            "缺失值不以0、均值或主观推断填充。"
        ),
    }
    (RUN_ROOT / "source_usage_audit.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({
        "normalized_envelopes": len(list(NORMALIZED.glob("*.json"))),
        "source_count": summary["source_count"],
        "candidate_records_before_merge": summary["candidate_records_before_cross_source_merge"],
        "audit": str(RUN_ROOT / "source_usage_audit.json"),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    sys.exit(main())
