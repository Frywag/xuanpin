from __future__ import annotations

import hashlib
import json
import re
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[2]
RUN_ROOT = Path(__file__).resolve().parent
OUTPUT = RUN_ROOT / "output"
DELIVERY = RUN_ROOT / "delivery"
CANONICAL_DOC = ROOT / "01.分级标准参考/01.统一业务范本/附录E_Canonical字段映射规格.md"
BASE_WORKBOOK = OUTPUT / "选品推荐表.xlsx"
DELIVERY_WORKBOOK = DELIVERY / "产品推荐表_全量数据_含入库152键_20260710.xlsx"
DELIVERY_JSONL = DELIVERY / "产品推荐_入库152键_20260710.jsonl"
DELIVERY_REPORT = DELIVERY / "数据源充分性与全量运行结论_20260710.md"

GRADE_FILL = {
    "S": "1F4E78",
    "A": "548235",
    "B": "BF8F00",
    "C": "7F6000",
}

SITE_SCOPE_KEYS = {
    "cp.listing_action", "cp.parentage", "cp.variation_theme", "cp.title_highlight",
    "cp.keywords", "cp.swatch_image", "cp.sku_image", "cp.offer_images",
    "cp.main_attribute", "cp.secondary_attribute", "cp.size_country",
    "cp.sale_price", "cp.msrp", "cp.price_minmax", "cp.b2b_pricing",
    "cp.tax_code", "cp.quantity", "cp.warehouse", "cp.handling_time",
    "cp.restock_date", "cp.inventory_always", "cp.max_order_qty", "cp.shelf_way",
    "cp.shipping_template", "cp.sku_classification", "cp.independent_packaging",
    "cp.location", "cp.import_designation", "cp.condition", "cp.gift_options",
    "cp.skip_offer", "cp.government_contract", "cp.collection_item",
}

SUPPLIER_KEYS = {
    "cp.gtin", "cp.model", "cp.unspsc", "cp.fabric_weight", "cp.weave_type",
    "cp.base_fabric", "cp.lining_composition", "cp.lining_weight",
    "cp.measure_bust", "cp.measure_waist", "cp.leg_opening", "cp.package_dims",
    "cp.package_weight", "cp.item_weight", "cp.country_of_origin",
    "cp.regulatory_cert", "cp.compliance_media", "cp.gpsr", "cp.pfas",
    "cp.baa_taa", "cp.certificate", "cp.green_purchasing", "cp.oem_sourced",
}

APPAREL_NA_KEYS = {
    "cp.liquid_capacity", "cp.blade_specs", "cp.battery_capacity",
    "cp.battery_compliance", "cp.dangerous_goods", "cp.ghs", "cp.hazmat",
    "cp.sds", "cp.pesticide", "cp.fcc", "cp.plated_metal",
}


def load_canonical_spec():
    domain = ""
    rows = []
    line_pattern = re.compile(
        r"^\| `(?P<key>cp\.[a-z0-9_]+)` \| (?P<variability>[^|]+) \| "
        r"(?P<meaning>[^|]+) \| (?P<amazon>[^|]+) \| (?P<temu>[^|]+) \| "
        r"(?P<shein>[^|]+) \|$")
    for line in CANONICAL_DOC.read_text(encoding="utf-8").splitlines():
        heading = re.match(r"^### (D\d+ .+)$", line)
        if heading:
            domain = heading.group(1)
            continue
        match = line_pattern.match(line)
        if match:
            item = {key: value.strip() for key, value in match.groupdict().items()}
            item["domain"] = domain
            rows.append(item)
    keys = [row["key"] for row in rows]
    if len(rows) != 152 or len(set(keys)) != 152:
        raise ValueError(f"Canonical 规格解析异常：rows={len(rows)} unique={len(set(keys))}")
    return rows


def present(value):
    return value not in (None, "", [], {})


def value_text(value):
    if value is None:
        return ""
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    return value


def normalize_label(value):
    return re.sub(r"[^a-z0-9]+", "", str(value or "").lower())


def json_objects(value):
    if isinstance(value, dict):
        return [value]
    if isinstance(value, list):
        return [item for item in value if isinstance(item, dict)]
    if not isinstance(value, str):
        return []
    objects = []
    try:
        parsed = json.loads(value)
        if isinstance(parsed, dict):
            return [parsed]
    except (ValueError, TypeError):
        pass
    decoder = json.JSONDecoder()
    for match in re.finditer(r"\{", value):
        try:
            parsed, _ = decoder.raw_decode(value[match.start():])
        except ValueError:
            continue
        if isinstance(parsed, dict):
            objects.append(parsed)
            break
    return objects


def flatten_structured(value, output):
    if isinstance(value, dict):
        for key, nested in value.items():
            if isinstance(nested, (dict, list)):
                flatten_structured(nested, output)
            elif present(nested):
                output[normalize_label(key)].append(nested)
    elif isinstance(value, list):
        for nested in value:
            flatten_structured(nested, output)


def record_field(records, names):
    for name in names:
        for record in records:
            value = record.get("fields", {}).get(name)
            if present(value):
                ev = record.get("evidence", {}).get(name) or {}
                evidence_ids = [ev.get("evidence_id")] if ev.get("evidence_id") else []
                return value, evidence_ids, record.get("source_id")
    return None, [], None


def record_evidence(records, names):
    evidence_ids = []
    for record in records:
        for name in names:
            ev = record.get("evidence", {}).get(name) or {}
            if ev.get("evidence_id"):
                evidence_ids.append(ev["evidence_id"])
    return sorted(set(evidence_ids))


def packet_evidence(packet, path):
    return sorted({
        item["evidence_id"] for item in packet.get("evidence_pack", [])
        if item.get("field_path") == path and item.get("evidence_id")
    })


def structured_index(records):
    output = defaultdict(list)
    for record in records:
        for field in ("product_attributes_json", "product_details_text",
                      "seller_sprite_selected_raw"):
            for obj in json_objects(record.get("fields", {}).get(field)):
                flatten_structured(obj, output)
    return output


def structured_value(index, aliases):
    for alias in aliases:
        values = index.get(normalize_label(alias), [])
        if values:
            return values[0]
    return None


def normalize_neckline(value):
    if not present(value):
        return None
    text = str(value).lower().replace("-", " ")
    for needle, label in (
        ("v neck", "V Neck"), ("crew neck", "Crew Neck"),
        ("scoop neck", "Scoop Neck"), ("square neck", "Square Neck"),
        ("halter", "Halter"), ("one shoulder", "One Shoulder"),
        ("collared", "Collared"),
    ):
        if needle in text:
            return label
    return str(value).strip()


def normalize_sleeve(value):
    if not present(value):
        return None
    text = str(value).lower().replace("-", " ")
    for needle, label in (
        ("sleeveless", "Sleeveless"), ("short sleeve", "Short Sleeve"),
        ("long sleeve", "Long Sleeve"), ("cap sleeve", "Cap Sleeve"),
        ("three quarter", "3/4 Sleeve"), ("3/4", "3/4 Sleeve"),
    ):
        if needle in text:
            return label
    return str(value).strip()


def normalize_silhouette(value):
    if not present(value):
        return None
    text = str(value).lower().replace("-", " ")
    for needle, label in (
        ("a line", "A-line"), ("bodycon", "Bodycon"),
        ("straight", "Straight"), ("fit and flare", "Fit and Flare"),
        ("flare", "Flare"),
    ):
        if needle in text:
            return label
    return str(value).strip()


def urls_from(value):
    if not present(value):
        return []
    if isinstance(value, (dict, list)):
        text = json.dumps(value, ensure_ascii=False)
    else:
        text = str(value)
    return list(dict.fromkeys(re.findall(r"https?://[^\s\"'<>\\]+", text)))


def build_canonical_152(packet, result, records, spec):
    values = {item["key"]: None for item in spec}
    statuses = {}
    evidence = {item["key"]: [] for item in spec}
    modes = {item["key"]: "" for item in spec}
    reasons = {}
    for item in spec:
        key = item["key"]
        if key in APPAREL_NA_KEYS:
            statuses[key] = "NOT_APPLICABLE_APPAREL_SCOPE"
            reasons[key] = "女装连衣裙选品范围内不适用；若后续产品含电气/液体/刀具特征需复核"
        elif key in SITE_SCOPE_KEYS:
            statuses[key] = "REQUIRES_SITE_SCOPE"
            reasons[key] = "需确定目标平台、站点、店铺和刊登动作后补齐"
        elif key in SUPPLIER_KEYS:
            statuses[key] = "REQUIRES_SUPPLIER_OR_COMPLIANCE_INPUT"
            reasons[key] = "现有市场/竞品数据不能替代供应商或合规事实"
        else:
            statuses[key] = "MISSING_SOURCE"
            reasons[key] = "本轮全量数据未提供可证据化取值"

    def assign(key, value, status, ev=None, mode="", reason=""):
        if key not in values or not present(value):
            return
        values[key] = value
        statuses[key] = status
        evidence[key] = sorted(set(ev or []))
        modes[key] = mode
        reasons[key] = reason

    basic = packet.get("basic_facts") or {}
    context = packet.get("context") or {}
    owned = packet.get("owned_supply_inputs") or {}
    structured = structured_index(records)
    semantic_evidence = record_evidence(
        records, ["product_attributes_json", "product_details_text", "description_text",
                  "material_text", "design_signals", "color_text", "size_text"])
    semantic_evidence = sorted(set(
        semantic_evidence + packet_evidence(packet, "basic_facts.title")))

    assign("cp.canonical_id", "CAN-" + hashlib.sha1(
        packet["candidate_id"].encode()).hexdigest()[:20].upper(),
        "FILLED_SYSTEM", [], "system", "由 candidate_id 确定性生成，重跑稳定")
    parent_asin, parent_ev, _ = record_field(records, ["parent_asin"])
    assign("cp.spu_id", parent_asin or packet.get("primary_product_id"),
           "FILLED_SOURCE_IDENTIFIER", parent_ev, "source_identifier")
    category = context.get("category_name") or (
        (basic.get("category_path") or [None])[-1]) or context.get("category_label")
    assign("cp.product_type", category, "FILLED_DIRECT_SOURCE",
           packet_evidence(packet, "basic_facts.category_path") +
           record_evidence(records, ["category_name"]), "direct")

    title_cn, title_cn_ev, _ = record_field(records, ["title_cn"])
    title_map = {}
    if basic.get("title"):
        title_map["en"] = basic["title"]
    if title_cn:
        title_map["zh"] = title_cn
    assign("cp.title_base", title_map, "FILLED_DIRECT_SOURCE",
           packet_evidence(packet, "basic_facts.title") + title_cn_ev, "direct")
    description, description_ev, _ = record_field(
        records, ["description_text", "product_details_text"])
    description_cn, description_cn_ev, _ = record_field(records, ["description_text_cn"])
    description_map = {}
    if description:
        description_map["en"] = str(description)
    if description_cn:
        description_map["zh"] = str(description_cn)
    assign("cp.description", description_map, "FILLED_DIRECT_SOURCE",
           description_ev + description_cn_ev, "direct")
    bullets = [part.strip() for part in str(description or "").split("|||") if part.strip()]
    if len(bullets) >= 2:
        assign("cp.bullet_points", bullets, "FILLED_DIRECT_SOURCE",
               description_ev, "direct")

    assign("cp.brand", basic.get("brand"), "FILLED_DIRECT_SOURCE",
           packet_evidence(packet, "basic_facts.brand"), "direct")
    assign("cp.main_image", basic.get("image_url"), "FILLED_DIRECT_SOURCE",
           packet_evidence(packet, "basic_facts.image_url"), "direct")
    selected_raw, selected_raw_ev, _ = record_field(records, ["seller_sprite_selected_raw"])
    raw_objects = json_objects(selected_raw)
    gallery_urls = []
    video_urls = []
    for raw_object in raw_objects:
        gallery_urls += urls_from(raw_object.get("allImages"))
        video_urls += urls_from(raw_object.get("video"))
    if basic.get("image_url") in gallery_urls:
        gallery_urls.remove(basic["image_url"])
    assign("cp.gallery_images", list(dict.fromkeys(gallery_urls)),
           "FILLED_STRUCTURED_SOURCE", selected_raw_ev, "structured")
    assign("cp.main_video", video_urls[0] if video_urls else None,
           "FILLED_STRUCTURED_SOURCE", selected_raw_ev, "structured")

    material_text, material_ev, _ = record_field(records, ["material_text"])
    material_type = structured_value(structured, ["Material Type", "Material", "Fabric Type"])
    fabric_type = structured_value(structured, ["Fabric Type"])
    composition_match = re.findall(
        r"\b\d{1,3}\s*%\s*[A-Za-z]+(?:\s+[A-Za-z]+)?", str(material_text or ""))
    assign("cp.composition", composition_match or material_text,
           "FILLED_STRUCTURED_SOURCE", material_ev or semantic_evidence, "structured")
    assign("cp.material_main", material_type or material_text,
           "FILLED_STRUCTURED_SOURCE", material_ev or semantic_evidence, "structured")
    assign("cp.fabric_type", fabric_type, "FILLED_STRUCTURED_SOURCE",
           semantic_evidence, "structured")
    assign("cp.fabric_weight", structured_value(
        structured, ["Apparel Fabric Weight Class"]), "FILLED_STRUCTURED_SOURCE",
        semantic_evidence, "structured")
    assign("cp.fabric_stretch", structured_value(
        structured, ["Apparel Fabric Stretch", "Stretch"]),
        "FILLED_STRUCTURED_SOURCE", semantic_evidence, "structured")
    assign("cp.lining_flag", structured_value(
        structured, ["Lining", "Lining Description"]),
        "FILLED_STRUCTURED_SOURCE", semantic_evidence, "structured")
    assign("cp.lining_desc", structured_value(
        structured, ["Lining Description"]), "FILLED_STRUCTURED_SOURCE",
        semantic_evidence, "structured")

    text = " | ".join(str(value) for value in (
        basic.get("title"), description, material_text,
        record_field(records, ["product_details_text"])[0]) if present(value))
    style = structured_value(structured, ["Apparel Occasion and Lifestyle", "Style"])
    pattern = structured_value(structured, ["Pattern"])
    season = structured_value(structured, ["Seasons", "Season"])
    neckline = normalize_neckline(structured_value(
        structured, ["Neck Style", "collar-type"]))
    sleeve_raw = structured_value(
        structured, ["Sleeve Length Description", "Sleeve Type"])
    sleeve_length = normalize_sleeve(sleeve_raw)
    silhouette = normalize_silhouette(structured_value(
        structured, ["Apparel Silhouette", "Silhouette"]))
    if not neckline:
        neckline = normalize_neckline(text)
    if not sleeve_length:
        sleeve_length = normalize_sleeve(text)
    if not silhouette:
        silhouette = normalize_silhouette(text)
    assign("cp.style", style, "FILLED_STRUCTURED_SOURCE", semantic_evidence, "structured")
    assign("cp.pattern", pattern, "FILLED_STRUCTURED_SOURCE", semantic_evidence, "structured")
    assign("cp.season", season, "FILLED_STRUCTURED_SOURCE", semantic_evidence, "structured")
    assign("cp.neckline", neckline, "FILLED_STRUCTURED_SOURCE", semantic_evidence, "structured")
    assign("cp.sleeve_length", sleeve_length, "FILLED_STRUCTURED_SOURCE",
           semantic_evidence, "structured")
    assign("cp.sleeve_type", structured_value(structured, ["Sleeve Type"]),
           "FILLED_STRUCTURED_SOURCE", semantic_evidence, "structured")
    assign("cp.closure", structured_value(structured, ["Apparel Closure Type"]),
           "FILLED_STRUCTURED_SOURCE", semantic_evidence, "structured")
    assign("cp.fit_type", structured_value(structured, ["Fit Type"]),
           "FILLED_STRUCTURED_SOURCE", semantic_evidence, "structured")
    assign("cp.silhouette", silhouette, "FILLED_STRUCTURED_SOURCE",
           semantic_evidence, "structured")
    assign("cp.length_style", structured_value(
        structured, ["Item Length Description"]), "FILLED_STRUCTURED_SOURCE",
        semantic_evidence, "structured")
    assign("cp.waistband", structured_value(structured, ["Waist Style"]),
           "FILLED_STRUCTURED_SOURCE", semantic_evidence, "structured")
    details = [structured_value(structured, [name]) for name in (
        "Additional Features", "Pocket Description")]
    assign("cp.detail_feature", [item for item in details if present(item)],
           "FILLED_STRUCTURED_SOURCE", semantic_evidence, "structured")
    care_text, care_ev, _ = record_field(records, ["care_text"])
    care = structured_value(structured, ["Product Care Instructions"]) or care_text
    assign("cp.care_instructions", care, "FILLED_STRUCTURED_SOURCE",
           semantic_evidence + care_ev, "structured")
    assign("cp.occasion", structured_value(
        structured, ["Apparel Occasion and Lifestyle", "Occasion Type"]),
        "FILLED_STRUCTURED_SOURCE", semantic_evidence, "structured")
    assign("cp.pocket", structured_value(structured, ["Pocket Description"]),
           "FILLED_STRUCTURED_SOURCE", semantic_evidence, "structured")
    assign("cp.back_style", structured_value(structured, ["Back Style"]),
           "FILLED_STRUCTURED_SOURCE", semantic_evidence, "structured")
    assign("cp.set_count", structured_value(
        structured, ["Unit Count", "Number of Items"]), "FILLED_STRUCTURED_SOURCE",
        semantic_evidence, "structured")

    scope_evidence = (packet_evidence(packet, "basic_facts.title") +
                      packet_evidence(packet, "basic_facts.category_path"))
    assign("cp.target_gender", "Women", "FILLED_RULE_SCOPE", scope_evidence,
           "scope_rule", "数据目录和类目均限定女装连衣裙")
    assign("cp.department", "Women", "FILLED_RULE_SCOPE", scope_evidence,
           "scope_rule", "数据目录和类目均限定女装连衣裙")
    assign("cp.age_range", structured_value(
        structured, ["Age Range Description"]), "FILLED_STRUCTURED_SOURCE",
        semantic_evidence, "structured")
    color_text, color_ev, _ = record_field(records, ["color_text"])
    assign("cp.color", structured_value(structured, ["Color"]) or color_text,
           "FILLED_STRUCTURED_SOURCE", semantic_evidence + color_ev, "structured")
    size_text, size_ev, _ = record_field(records, ["size_text"])
    assign("cp.size_value", size_text, "FILLED_DIRECT_SOURCE", size_ev, "direct")
    assign("cp.size_system", structured_value(
        structured, ["Garment Size Country"]), "FILLED_STRUCTURED_SOURCE",
        semantic_evidence, "structured")
    assign("cp.special_size", structured_value(
        structured, ["Special Size Type"]), "FILLED_STRUCTURED_SOURCE",
        semantic_evidence, "structured")
    assign("cp.measure_length", structured_value(
        structured, ["Shoulder to Bottom Hem Length"]), "FILLED_STRUCTURED_SOURCE",
        semantic_evidence, "structured")
    assign("cp.fit_sentiment", structured_value(
        structured, ["Fit to Size Sentiment"]), "FILLED_STRUCTURED_SOURCE",
        semantic_evidence, "structured")

    original_price = basic.get("original_price") or {}
    price = basic.get("price") or {}
    assign("cp.list_price", original_price.get("amount"), "FILLED_DIRECT_SOURCE",
           packet_evidence(packet, "basic_facts.original_price"), "direct")
    assign("cp.your_price", price.get("amount"), "FILLED_DIRECT_SOURCE",
           packet_evidence(packet, "basic_facts.price"), "direct")
    assign("cp.currency", price.get("currency"), "FILLED_DIRECT_SOURCE",
           packet_evidence(packet, "basic_facts.price"), "direct")
    canonical_url, canonical_ev, _ = record_field(records, ["canonical_url"])
    assign("cp.reference_link", packet.get("canonical_url") or canonical_url,
           "FILLED_DIRECT_SOURCE", canonical_ev, "direct")
    fulfillment, fulfillment_ev, _ = record_field(records, ["seller_fulfillment"])
    assign("cp.fulfillment_channel", fulfillment, "FILLED_DIRECT_SOURCE",
           fulfillment_ev, "direct")
    launch_date, launch_ev, _ = record_field(records, ["first_available_date"])
    assign("cp.launch_date", launch_date, "FILLED_DIRECT_SOURCE", launch_ev, "direct")
    for key, field_name in (
        ("cp.package_dims", "package_dimensions"),
        ("cp.package_weight", "package_weight"),
        ("cp.item_weight", "item_weight"),
        ("cp.package_type", "package_type"),
    ):
        raw_value, raw_ev, _ = record_field(records, [field_name])
        assign(key, raw_value, "FILLED_DIRECT_SOURCE", raw_ev, "direct")
    assign("cp.number_of_items", structured_value(
        structured, ["Number of Items"]), "FILLED_STRUCTURED_SOURCE",
        semantic_evidence, "structured")
    assign("cp.ships_globally", owned.get("ships_globally"),
           "FILLED_DIRECT_SOURCE", [], "owned")

    filled = sum(present(value) for value in values.values())
    return {
        "values": values,
        "statuses": statuses,
        "evidence": evidence,
        "modes": modes,
        "reasons": reasons,
        "summary": {
            "key_count": len(values),
            "filled_count": filled,
            "filled_pct": round(filled / len(values) * 100, 1),
            "status_counts": dict(Counter(statuses.values())),
        },
    }


def load_run_data():
    packets = {}
    with (OUTPUT / "candidate_packets.jsonl").open(encoding="utf-8") as handle:
        for line in handle:
            packet = json.loads(line)
            packets[packet["candidate_id"]] = packet
    results = {}
    with (OUTPUT / "analysis_results.jsonl").open(encoding="utf-8") as handle:
        for line in handle:
            result = json.loads(line)
            results[result["candidate_id"]] = result
    record_index = defaultdict(list)
    for path in sorted((OUTPUT / "envelopes").glob("*.json")):
        envelope = json.loads(path.read_text(encoding="utf-8"))
        for record in envelope.get("records") or []:
            if record.get("record_kind") != "product":
                continue
            platform = record.get("platform") or record.get("site") or "tiktok_shop"
            product_id = str(record.get("product_id") or "")
            if not product_id:
                continue
            copy = dict(record)
            copy["source_id"] = envelope["source_id"]
            copy["source_type"] = envelope["source_type"]
            record_index[(platform, product_id)].append(copy)
    return packets, results, record_index


def delivery_records():
    spec = load_canonical_spec()
    packets, results, record_index = load_run_data()
    ordered = sorted(results.values(), key=lambda item: -(item["priority_score"].get("pct") or 0))
    deliveries = []
    for rank, result in enumerate(ordered, 1):
        if result["priority_score"].get("grade") == "C":
            continue
        packet = packets[result["candidate_id"]]
        records = record_index.get((packet["platform"], str(packet.get("primary_product_id") or "")), [])
        canonical = build_canonical_152(packet, result, records, spec)
        deliveries.append({
            "schema_version": "canonical.product.selection.v1",
            "rank": rank,
            "candidate_id": packet["candidate_id"],
            "grade": result["priority_score"].get("grade"),
            "pct": result["priority_score"].get("pct"),
            "priority_score": result["priority_score"].get("total"),
            "track": result["priority_score"].get("track_label") or result["priority_score"].get("track"),
            "platform": packet["platform"],
            "title": packet.get("basic_facts", {}).get("title"),
            "product_link": packet.get("canonical_url"),
            "source_refs": packet.get("source_refs") or [],
            "confidence": result.get("confidence"),
            "recommendation": result.get("recommendation") or {},
            "canonical_key_count": 152,
            "canonical_152": canonical["values"],
            "canonical_status": canonical["statuses"],
            "canonical_evidence": canonical["evidence"],
            "canonical_fill_mode": canonical["modes"],
            "canonical_missing_reason": canonical["reasons"],
            "canonical_summary": canonical["summary"],
        })
    return spec, packets, results, deliveries


def style_header(ws, row=1):
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[row]:
        cell.fill = fill
        cell.font = Font(name="Arial", color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def set_widths(ws, widths):
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width


def add_simple_sheet(workbook, name, headers, rows, widths, index=None):
    if name in workbook.sheetnames:
        del workbook[name]
    ws = workbook.create_sheet(name, index)
    ws.append(headers)
    for row in rows:
        ws.append([value_text(value) for value in row])
    style_header(ws)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    set_widths(ws, widths)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    return ws


def build_workbook(spec, packets, results, deliveries, audit, run_meta):
    workbook = load_workbook(BASE_WORKBOOK)
    if "推荐总表" in workbook.sheetnames:
        workbook["推荐总表"].title = "全量分级底稿"
    if "商品画像" in workbook.sheetnames:
        workbook["商品画像"].title = "辅助画像标签"
    for name in ("交付说明", "产品推荐", "入库152键_宽表", "入库152键_长表",
                 "152键字典", "数据源使用审计"):
        if name in workbook.sheetnames:
            del workbook[name]

    explanation_rows = [
        ("交付结论", "数据足够用于本轮女装连衣裙选品推荐与 Canonical 152 键占位入库；不等于已满足平台刊登。"),
        ("全量候选", run_meta["candidates_total"]),
        ("等级分布", run_meta["grade_distribution"]),
        ("正式推荐/备选", f"S/A/B 共 {len(deliveries)} 款；其中 S/A {sum(item['grade'] in ('S','A') for item in deliveries)} 款"),
        ("数据源", f"25 个来源：SellerSprite、Kalodata、23 个前端站点工作簿"),
        ("全量使用口径", audit["full_use_policy"]),
        ("噪声排除", "Revolve 500 行全部无商品ID/URL，抽查为排序/查看/收藏控件文本；已审计但不进入 L1。"),
        ("152 键口径", "严格取附录 E 的 152 个唯一 cp.* 键；宽表每行固定 152 键，长表逐键给状态与证据。"),
        ("缺失策略", "缺失不填 0、不补均值、不主观猜测；SiteScope、供应链和合规事实分别标记待补。"),
        ("评分性质", "L1/L2/L3 与 S/A/B/C 均由确定性引擎产生；本交付未用 LLM 改分。"),
        ("机器入库文件", DELIVERY_JSONL.name),
        ("全量底稿", "工作簿中的‘全量分级底稿’保留全部 1,800 个候选，含 C 级淘汰依据。"),
    ]
    ws = add_simple_sheet(workbook, "交付说明", ["项目", "内容"], explanation_rows,
                          [28, 110], 0)
    ws.sheet_view.showGridLines = False

    recommendation_rows = []
    for item in deliveries:
        packet = packets[item["candidate_id"]]
        result = results[item["candidate_id"]]
        price = packet["basic_facts"].get("price") or {}
        recommendation = item["recommendation"]
        recommendation_rows.append([
            item["rank"], item["grade"], item["pct"],
            recommendation.get("decision"), item["track"], item["platform"],
            item["candidate_id"], item["title"], packet["basic_facts"].get("brand"),
            price.get("amount"), price.get("currency"),
            packet["market_metrics"].get("sales_30d_units"),
            packet["basic_facts"].get("rating"), packet["basic_facts"].get("review_count"),
            "、".join(ref["source_id"] for ref in item["source_refs"]), item["confidence"],
            recommendation.get("reason_summary"), recommendation.get("risk_summary"),
            "；".join(recommendation.get("next_actions") or []),
            item["canonical_summary"]["filled_count"],
            item["canonical_summary"]["filled_pct"], item["product_link"],
        ])
    ws = add_simple_sheet(
        workbook, "产品推荐",
        ["排名", "等级", "pct", "决策", "赛道", "平台", "candidate_id", "产品标题",
         "品牌", "价格", "币种", "近30天销量", "评分", "评论数", "数据源", "置信度",
         "推荐理由", "风险摘要", "下一步", "152键已填数", "152键覆盖率%", "产品链接"],
        recommendation_rows,
        [8, 7, 8, 10, 14, 14, 28, 48, 18, 10, 8, 13, 8, 10, 32, 10,
         58, 48, 48, 13, 14, 42], 1)
    for row in range(2, ws.max_row + 1):
        grade_cell = ws.cell(row, 2)
        if grade_cell.value in GRADE_FILL:
            grade_cell.fill = PatternFill("solid", fgColor=GRADE_FILL[grade_cell.value])
            grade_cell.font = Font(name="Arial", color="FFFFFF", bold=True)
        link_cell = ws.cell(row, 22)
        if link_cell.value:
            link_cell.hyperlink = str(link_cell.value)
            link_cell.style = "Hyperlink"

    keys = [item["key"] for item in spec]
    wide_headers = ["排名", "等级", "pct", "平台", "candidate_id", "产品标题",
                    "152键已填数", "152键覆盖率%"] + keys
    wide_rows = [[
        item["rank"], item["grade"], item["pct"], item["platform"],
        item["candidate_id"], item["title"], item["canonical_summary"]["filled_count"],
        item["canonical_summary"]["filled_pct"],
        *[item["canonical_152"][key] for key in keys],
    ] for item in deliveries]
    ws = add_simple_sheet(workbook, "入库152键_宽表", wide_headers, wide_rows,
                          [8, 7, 8, 14, 28, 45, 13, 14] + [18] * 152, 2)
    ws.freeze_panes = "I2"

    spec_by_key = {item["key"]: item for item in spec}
    long_rows = []
    for item in deliveries:
        for key in keys:
            meta = spec_by_key[key]
            long_rows.append([
                item["rank"], item["grade"], item["pct"], item["platform"],
                item["candidate_id"], meta["domain"], key, meta["variability"],
                meta["meaning"], item["canonical_152"][key],
                item["canonical_status"][key], item["canonical_fill_mode"][key],
                ",".join(item["canonical_evidence"][key]),
                item["canonical_missing_reason"][key],
            ])
    add_simple_sheet(
        workbook, "入库152键_长表",
        ["排名", "等级", "pct", "平台", "candidate_id", "域", "canonical_key", "变/不变",
         "业务含义", "值(JSON)", "状态", "填充方式", "evidence_refs", "缺失/状态说明"],
        long_rows, [8, 7, 8, 14, 28, 18, 28, 10, 28, 45, 34, 18, 44, 55], 3)

    dictionary_rows = [[
        item["domain"], item["key"], item["variability"], item["meaning"],
        item["amazon"], item["temu"], item["shein"],
    ] for item in spec]
    add_simple_sheet(
        workbook, "152键字典",
        ["域", "canonical_key", "变/不变", "业务含义", "Amazon映射", "Temu映射", "Shein映射"],
        dictionary_rows, [18, 28, 10, 28, 52, 42, 42], 4)

    audit_rows = []
    for source in audit["sources"]:
        audit_rows.append([
            source["source_id"], source["source_type"], source["artifact_path"],
            source.get("candidate_records", 0), source.get("partial_records", 0),
            source.get("skipped_records", 0), source.get("sheet_rows_scanned") or
            source.get("jsonl_rows_available"), source.get("package_files_verified"),
            source.get("usage"), "；".join(source.get("warnings") or []),
        ])
    add_simple_sheet(
        workbook, "数据源使用审计",
        ["source_id", "类型", "原始产物", "有效候选记录", "降级记录", "审计排除记录",
         "扫描/可用行数", "包文件数", "实际用途", "边界/警告"],
        audit_rows, [28, 12, 62, 14, 12, 14, 55, 12, 65, 75], 5)

    for sheet in workbook.worksheets:
        sheet.sheet_view.showGridLines = False
        for row in sheet.iter_rows():
            for cell in row:
                if not cell.font.name or cell.font.name == "Calibri":
                    cell.font = Font(
                        name="Arial", size=cell.font.sz or 10, bold=cell.font.bold,
                        italic=cell.font.italic, color=cell.font.color)
    workbook.save(DELIVERY_WORKBOOK)


def write_report(deliveries, audit, run_meta):
    grade_dist = run_meta["grade_distribution"]
    selected_dist = Counter(item["grade"] for item in deliveries)
    filled_counts = [item["canonical_summary"]["filled_count"] for item in deliveries]
    status_totals = Counter()
    for item in deliveries:
        status_totals.update(item["canonical_status"].values())
    lines = [
        "# 06.分级系统数据源充分性与全量运行结论",
        "",
        "## 结论",
        "",
        "**足够用于本轮女装连衣裙的选品推荐、S/A/B/C 分级和 Canonical 152 键占位入库；不等于已具备直接铺货/刊登条件。**",
        "",
        "本轮全量输入覆盖 SellerSprite 的 Amazon 需求/竞争/趋势数据、Kalodata 的 TikTok Shop 市场与商品数据，以及 23 个前端站点工作簿。供应链成本、MOQ、交期、原产地和部分平台合规字段仍需在终选后由供应商、合规或 SiteScope 流程补齐。",
        "",
        "## 实跑结果",
        "",
        f"- 唯一候选：{run_meta['candidates_total']:,} 款。",
        f"- 等级分布：S={grade_dist.get('S', 0)}、A={grade_dist.get('A', 0)}、B={grade_dist.get('B', 0)}、C={grade_dist.get('C', 0)}。",
        f"- 正式推荐与观察备选：S/A/B 共 {len(deliveries):,} 款，其中 S/A 共 {selected_dist['S'] + selected_dist['A']:,} 款。",
        f"- L2/L3：{run_meta['l2_processed']} / {run_meta['l3_processed']} 款，符合成本上限。",
        f"- 152 键：每个 S/A/B 产品固定输出 152 个唯一 `cp.*` 键；已填键数范围 {min(filled_counts)}–{max(filled_counts)}，其余逐键标记缺失原因。",
        "",
        "## 全量使用边界",
        "",
        f"- 规范化前有效候选记录：{audit['candidate_records_before_cross_source_merge']:,}；同平台同商品跨源合并后为 {run_meta['candidates_total']:,} 个唯一候选。",
        "- SellerSprite：200 款产品总表全部进入评分；销量、BSR、价格、评论趋势全部有效行聚合；全字段长表逐行扫描并保留字段目录与关键原值。",
        "- Kalodata：93 个标准商品记录全部进入评分；无损包的 raw/processed/ledger 文件继续作为审计底稿。",
        "- 前端：所有工作簿均被扫描；有效商品进入评分或详情增强。Revolve 的 500 行无 ID/URL，经抽查为排序、查看数量和收藏控件文本，已审计排除。",
        "- 未映射进 Priority Score 的原始字段没有删除，仍通过标准信封和原始文件路径保留，供 152 键或后续复核使用。",
        "",
        "## 仍需补齐",
        "",
        "- 供应链：采购成本、MOQ、交期、包装实测、质检结果。",
        "- 合规：原产地、检测/证书、GPSR/PFAS/Prop65 等目标平台要求。",
        "- SiteScope：目标平台、站点、店铺、库存、仓、促销价、刊登动作和变体展开。",
        "- 评论深挖：当前有评分/评论量和部分摘要，缺少全量 1–3 星评论正文时，L2 痛点结论需继续补采。",
        "",
        "## 交付文件",
        "",
        f"- `{DELIVERY_WORKBOOK.name}`：业务推荐表、全量分级底稿、152 键宽/长表、证据和源使用审计。",
        f"- `{DELIVERY_JSONL.name}`：机器入库 JSONL；每条含 `canonical_152`、`canonical_status` 和 `canonical_evidence`。",
        "- `delivery_manifest.json`：交付文件 SHA-256。",
        "",
        "## 状态统计",
        "",
    ]
    lines += [f"- `{status}`：{count:,}" for status, count in status_totals.most_common()]
    DELIVERY_REPORT.write_text("\n".join(lines) + "\n", encoding="utf-8")


def sha256(path):
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def main():
    DELIVERY.mkdir(parents=True, exist_ok=True)
    spec, packets, results, deliveries = delivery_records()
    with DELIVERY_JSONL.open("w", encoding="utf-8") as handle:
        for item in deliveries:
            handle.write(json.dumps(item, ensure_ascii=False, separators=(",", ":")) + "\n")
    audit = json.loads((RUN_ROOT / "source_usage_audit.json").read_text(encoding="utf-8"))
    run_meta = json.loads((OUTPUT / "run_meta.json").read_text(encoding="utf-8"))
    build_workbook(spec, packets, results, deliveries, audit, run_meta)
    write_report(deliveries, audit, run_meta)
    spec_path = DELIVERY / "canonical_152_spec.json"
    spec_path.write_text(json.dumps(spec, ensure_ascii=False, indent=2), encoding="utf-8")
    manifest = {
        "generated_from_run": run_meta["run_id"],
        "canonical_key_count": len(spec),
        "delivery_products": len(deliveries),
        "files": [],
    }
    for path in (DELIVERY_WORKBOOK, DELIVERY_JSONL, DELIVERY_REPORT, spec_path):
        manifest["files"].append({
            "path": path.name,
            "bytes": path.stat().st_size,
            "sha256": sha256(path),
        })
    (DELIVERY / "delivery_manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({
        "delivery_products": len(deliveries),
        "canonical_keys": len(spec),
        "workbook": str(DELIVERY_WORKBOOK),
        "jsonl": str(DELIVERY_JSONL),
        "report": str(DELIVERY_REPORT),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
