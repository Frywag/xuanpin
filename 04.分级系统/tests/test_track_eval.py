# -*- coding: utf-8 -*-
"""三赛道独立评估：验收场景测试（交接清单 §12）。"""
import copy
from pathlib import Path

import pytest
import yaml

from grading_system.models import CandidateDataPacket, EvidenceRef
from grading_system.scoring import GroupStats
from grading_system.track_eval import TRACK_IDS, TrackEvaluator, assess_freshness

CFG = yaml.safe_load((Path(__file__).resolve().parents[1] /
                      "configs/tracks_v1.yaml").read_text(encoding="utf-8"))


def strict_cfg():
    """关闭临时可跑占位规则后的严格证据口径（回归对照用）。"""
    c = copy.deepcopy(CFG)
    for t in c["tracks"].values():
        if "provisional_admission" in t:
            t["provisional_admission"]["enabled"] = False
    return c


def ev(path, eid, src="s1"):
    return EvidenceRef(eid, src, "frontend", "T", path, "a.xlsx", "sheet=s;row=2")


def base_packet(cid="c1", group="indie_frontend", roles=("discovery",)):
    p = CandidateDataPacket(cid, "US", "site", "p1", "https://x/p1")
    p.source_refs = [{"source_id": group, "envelope_id": "e1"}]
    p.context["source_group"] = group
    p.context["source_roles"] = list(roles)
    p.set_fact("basic_facts", "title", "Lace Maxi Dress", ev("basic_facts.title", f"ev_{cid}_t"))
    return p


def make_evaluator(packets):
    return TrackEvaluator(CFG, GroupStats(packets), "run_t")


class TestInvariants:
    def test_three_records_per_candidate(self):
        """场景/验收：每个候选每个 run 恰好三条候选级 TrackEvaluation。"""
        packets = [base_packet(f"c{i}") for i in range(5)]
        evals, _, _ = make_evaluator(packets).run(packets)
        assert len(evals) == 15
        for p in packets:
            assert {e.track_id for e in evals if e.subject_id == p.candidate_id} == set(TRACK_IDS)

    def test_pending_and_rejected_have_no_grade(self):
        p = base_packet()
        evals, _, _ = make_evaluator([p]).run([p])
        for e in evals:
            if e.admission_status != "ELIGIBLE":
                assert e.grade is None          # 正式 grade 为空
                assert e.legacy_grade == "C"    # 仅旧表兼容

    def test_determinism(self):
        packets = [base_packet(f"c{i}") for i in range(3)]
        r1 = make_evaluator(packets).run(packets)[0]
        packets2 = [base_packet(f"c{i}") for i in range(3)]
        r2 = make_evaluator(packets2).run(packets2)[0]
        assert [(e.subject_id, e.track_id, e.admission_status, e.grade) for e in r1] \
            == [(e.subject_id, e.track_id, e.admission_status, e.grade) for e in r2]

    def test_results_carry_calibration_status(self):
        p = base_packet()
        evals, _, _ = make_evaluator([p]).run([p])
        assert all("calibration" in e.rule_status for e in evals)


class TestTrendNew:
    def test_scene1_no_sales_but_both_evidences_is_eligible(self):
        """场景一：有近期推出+款型新颖度证据、无任何销量 -> ELIGIBLE 并给级。"""
        p = base_packet("cn1")
        p.freshness.update({"freshness_status": "confirmed",
                            "new_arrival_flag": "page_new_arrival",
                            "novelty_status": "confirmed",
                            "freshness_evidence_refs": ["ev_f1"],
                            "novelty_evidence_refs": ["ev_n1"]})
        e = make_evaluator([p]).eval_trend_new(p, None)
        assert e.admission_status == "ELIGIBLE"
        assert e.grade in ("S", "A", "B")   # 无销量绝不淘汰（单源早期可 B）
        assert e.evidence_refs              # 新款证据 100% 绑定

    def test_scene7_first_seen_only_admits_provisionally(self):
        """场景七更新（业务指示 2026-07-14）：无上架/新颖度证据属数据源结构性
        缺失 -> 按占位规则临时准入并照常评分；缺失照记、依据标注。"""
        p = base_packet("cn2")
        assess_freshness(p)
        e = make_evaluator([p]).eval_trend_new(p, None)
        assert e.admission_status == "ELIGIBLE"
        assert e.admission_basis == "provisional_rule"
        assert e.grade is not None
        assert e.confidence == "low"
        assert e.missing_fields          # 缺失仍然点出来
        assert e.provisional_notes

    def test_scene7_strict_config_still_pending(self):
        """关闭占位规则（enabled: false）即回到严格证据口径 -> PENDING_DATA。"""
        p = base_packet("cn2s")
        assess_freshness(p)
        e = TrackEvaluator(strict_cfg(), GroupStats([p]), "run_t").eval_trend_new(p, None)
        assert e.admission_status == "PENDING_DATA"
        assert e.grade is None

    def test_scene8_confirmed_old_style_rejected(self):
        p = base_packet("cn3")
        p.freshness["novelty_status"] = "old"
        p.market_metrics["sales_growth_rate"] = 1.5   # 销量增长不能覆盖旧款结论
        e = make_evaluator([p]).eval_trend_new(p, None)
        assert e.admission_status == "REJECTED"

    def test_scene5_indie_discovery_not_structurally_rejected(self):
        """场景五：独立站缺交易字段，按发现源参与趋势判断，不被淘汰为 REJECTED
        （占位规则下临时准入参赛；确证旧款仍会 REJECTED）。"""
        p = base_packet("cn4", group="indie_frontend")
        assess_freshness(p)
        e = make_evaluator([p]).eval_trend_new(p, None)
        assert e.admission_status != "REJECTED"
        assert e.admission_status == "ELIGIBLE"       # 临时准入，照常竞争

    def test_shein_new_arrival_collection_is_recent_evidence(self):
        p = base_packet("cn5", group="shein_frontend")
        p.set_fact("market_metrics", "bestseller_rank_text",
                   "#8 Bestseller in New Arrival Mini Casual Dresses",
                   ev("market_metrics.bestseller_rank_text", "ev_cn5_b"))
        assess_freshness(p)
        assert p.freshness["freshness_status"] == "inferred"
        assert p.freshness["freshness_evidence_refs"]


class TestHitImprovement:
    def _demand_packet(self, cid, rating, rc, clusters=0):
        p = base_packet(cid, group="shein_frontend", roles=("transaction", "voc"))
        p.set_fact("market_metrics", "sales_floor_units", 800,
                   ev("market_metrics.sales_floor_units", f"ev_{cid}_s"))
        if rating is not None:
            p.set_fact("basic_facts", "rating", rating, ev("basic_facts.rating", f"ev_{cid}_r"))
            p.set_fact("basic_facts", "review_count", rc,
                       ev("basic_facts.review_count", f"ev_{cid}_rc"))
        p.product_opportunity["negative_review_clusters"] = [
            {"cluster_name": f"痛点{i}", "mention_count": 5,
             "evidence_refs": [f"ev_{cid}_c{i}"]} for i in range(clusters)]
        return p

    def test_scene3_joint_admission(self):
        p = self._demand_packet("h1", rating=3.8, rc=120, clusters=3)
        e = make_evaluator([p]).eval_hit_improvement(p)
        assert e.admission_status == "ELIGIBLE"
        assert e.grade in ("S", "A", "B")

    def test_scene4_low_rating_small_sample_not_eligible(self):
        """场景四：低评分但样本很少 -> 不得进入高等级（待补）。"""
        p = self._demand_packet("h2", rating=3.6, rc=5, clusters=0)
        e = make_evaluator([p]).eval_hit_improvement(p)
        assert e.admission_status == "PENDING_DATA"

    def test_normal_rating_rejected(self):
        p = self._demand_packet("h3", rating=4.7, rc=200, clusters=0)
        e = make_evaluator([p]).eval_hit_improvement(p)
        assert e.admission_status == "REJECTED"

    def test_missing_voc_admits_provisionally_capped_b(self):
        """业务指示 2026-07-14：评论全文属源结构性缺失——有交易信号+确证低评分
        即占位豁免负面聚类准入；痛点未确认，占位等级上限 B。"""
        p = self._demand_packet("h4", rating=3.8, rc=120, clusters=0)
        e = make_evaluator([p]).eval_hit_improvement(p)
        assert e.admission_status == "ELIGIBLE"
        assert e.admission_basis == "provisional_rule"
        assert e.grade == "B"            # grade_cap: B（占位）
        assert any("聚类" in m or "评论" in m for m in e.missing_fields)

    def test_missing_voc_strict_config_still_pending(self):
        p = self._demand_packet("h4s", rating=3.8, rc=120, clusters=0)
        e = TrackEvaluator(strict_cfg(), GroupStats([p]), "run_t").eval_hit_improvement(p)
        assert e.admission_status == "PENDING_DATA"

    def test_no_rating_cannot_admit_provisionally(self):
        """低分前提无法确证（缺评分）时占位规则不适用 -> 仍 PENDING。"""
        p = self._demand_packet("h5", rating=None, rc=0, clusters=0)
        e = make_evaluator([p]).eval_hit_improvement(p)
        assert e.admission_status == "PENDING_DATA"


class TestLongTailAndBudgets:
    def test_not_a_catchall(self):
        """未命中前两赛道 ≠ 自动进入长尾：无交易证据 -> PENDING_DATA
        （稳定需求是本赛道前提，占位规则不豁免）。"""
        p = base_packet("l1")
        e = make_evaluator([p]).eval_long_tail(p)
        assert e.admission_status == "PENDING_DATA"

    def test_missing_price_admits_provisionally_capped_b(self):
        """业务指示 2026-07-14：仅缺价格（源结构性没有）时占位准入，
        价格/利润结构不可判 -> 占位等级上限 B。"""
        p = base_packet("l2", group="shein_frontend", roles=("transaction",))
        p.set_fact("market_metrics", "sales_floor_units", 400,
                   ev("market_metrics.sales_floor_units", "ev_l2_s"))
        e = make_evaluator([p]).eval_long_tail(p)
        assert e.admission_status == "ELIGIBLE"
        assert e.admission_basis == "provisional_rule"
        assert e.grade == "B"
        assert any(m.startswith("价格") for m in e.missing_fields)

    def test_spike_exclusion_not_waived_by_provisional(self):
        """单次暴涨排除是抗噪红线，占位规则不豁免 -> 仍 PENDING。"""
        p = base_packet("l3", group="shein_frontend", roles=("transaction",))
        p.set_fact("market_metrics", "sales_floor_units", 400,
                   ev("market_metrics.sales_floor_units", "ev_l3_s"))
        p.set_fact("market_metrics", "sales_growth_rate", 1.2,
                   ev("market_metrics.sales_growth_rate", "ev_l3_g"))
        p.set_fact("basic_facts", "price", {"amount": 19.9, "currency": "USD"},
                   ev("basic_facts.price", "ev_l3_p"))
        e = make_evaluator([p]).eval_long_tail(p)
        assert e.admission_status == "PENDING_DATA"

    def test_scene6_budgets_independent(self):
        """场景六：三赛道 L2/L3 名额互不挤占（各自独立队列）。"""
        packets = [base_packet(f"b{i}") for i in range(4)]
        packets[0].freshness.update({"freshness_status": "confirmed",
                                     "novelty_status": "confirmed"})
        _, _, budgets = make_evaluator(packets).run(packets)
        assert set(budgets) == set(TRACK_IDS)
        for tid in TRACK_IDS:
            assert set(budgets[tid]) == {"l2_queue", "l3_queue", "refetch_queue"}

    def test_refetch_budget_split_no_deadlock(self):
        """P0-03：补证预算与深挖预算拆分——缺证据候选（临时准入或 PENDING）
        进入 refetch_queue 获得本轮补采名额，不再「缺数据者永远无预算」。"""
        packets = [base_packet(f"p{i}") for i in range(6)]
        _, _, budgets = make_evaluator(packets).run(packets)
        q = budgets["trend_new"]["refetch_queue"]
        assert q, "缺证据候选必须获得补证名额"
        assert len(q) <= CFG["tracks"]["trend_new"]["refetch_budget_draft"]
        for item in q:
            assert set(item) == {"subject_id", "missing", "tasks",
                                 "status", "owner", "due", "basis"}
            assert item["status"] == "open"
            assert item["basis"] in ("provisional_eligible", "pending")
            assert item["tasks"], "补证队列条目必须带可执行补采任务"
            assert item["missing"], "补证队列条目必须点出缺失证据"

    def test_refetch_queue_strict_config_from_pending(self):
        """严格口径（占位规则关闭）下补证队列仍从 PENDING_DATA 取（P0-03 原语义）。"""
        packets = [base_packet(f"q{i}") for i in range(6)]
        _, _, budgets = TrackEvaluator(strict_cfg(), GroupStats(packets),
                                       "run_t").run(packets)
        q = budgets["trend_new"]["refetch_queue"]
        assert q and all(i["basis"] == "pending" for i in q)
        # 严格口径下补证队列（PENDING）与深挖队列（ELIGIBLE）互不重叠
        assert not ({i["subject_id"] for i in q}
                    & set(budgets["trend_new"]["l2_queue"]))

    def test_all_tracks_have_refetch_budget(self):
        """每条赛道都必须配置独立补证预算（tracks_v1.yaml）。"""
        for tid in TRACK_IDS:
            assert CFG["tracks"][tid].get("refetch_budget_draft", 0) > 0


class TestTrendFairness:
    """业务确认（2026-07-13）：趋势赛道不要求评论等数据点位，
    独立站与交易平台公平竞争；评论全量采集与分析后置 L3。"""

    def _fresh(self, p):
        p.freshness.update({"freshness_status": "confirmed",
                            "new_arrival_flag": "page_new_arrival",
                            "novelty_status": "confirmed",
                            "freshness_evidence_refs": ["ev_f"],
                            "novelty_evidence_refs": ["ev_n"]})
        return p

    def test_reviews_and_sales_do_not_affect_trend_result(self):
        """同一候选带不带评论/销量，趋势准入与分数必须完全一致。"""
        bare = self._fresh(base_packet("fa1", group="indie_frontend"))
        rich = self._fresh(base_packet("fa1", group="indie_frontend"))
        rich.set_fact("basic_facts", "rating", 4.9, ev("basic_facts.rating", "ev_r"))
        rich.set_fact("basic_facts", "review_count", 5000,
                      ev("basic_facts.review_count", "ev_rc"))
        rich.set_fact("market_metrics", "sales_30d_units", 99999,
                      ev("market_metrics.sales_30d_units", "ev_s"))
        e1 = make_evaluator([bare]).eval_trend_new(bare, None)
        e2 = make_evaluator([rich]).eval_trend_new(rich, None)
        assert (e1.admission_status, e1.grade, e1.score) \
            == (e2.admission_status, e2.grade, e2.score)

    def test_indie_can_reach_top_grades_without_reviews(self):
        """独立站无评论无销量，凭新款证据+设计细节+机会簇可公平竞争高等级。"""
        p = self._fresh(base_packet("fa2", group="indie_frontend"))
        p.context["design_signals"] = [{"kind": "design_details", "text": "wrap maxi"}]
        p.add_evidence(ev("context.design_signals", "ev_d"), "context.design_signals")
        e = make_evaluator([p]).eval_trend_new(p, cluster_id="clu_x")
        assert e.admission_status == "ELIGIBLE"
        assert e.grade in ("S", "A")   # 不因缺评论/销量被压制

    def test_review_collection_deferred_to_l3(self):
        """趋势赛道 L2 补采不含评论；评论出现在 L3 终选验证清单。"""
        tcfg = CFG["tracks"]["trend_new"]
        assert not any("评论" in x for x in tcfg["l2_focus"])
        assert any("评论" in x for x in tcfg["l3_focus"])

    def test_provisional_path_also_blind_to_reviews_and_sales(self):
        """占位规则准入路径同样不得读取评论/销量：带不带这些字段结果必须一致。"""
        bare = base_packet("fp1", group="indie_frontend")
        rich = base_packet("fp1", group="indie_frontend")
        rich.set_fact("basic_facts", "rating", 4.9, ev("basic_facts.rating", "ev_pr"))
        rich.set_fact("basic_facts", "review_count", 5000,
                      ev("basic_facts.review_count", "ev_prc"))
        rich.set_fact("market_metrics", "sales_30d_units", 99999,
                      ev("market_metrics.sales_30d_units", "ev_ps"))
        e1 = make_evaluator([bare]).eval_trend_new(bare, None)
        e2 = make_evaluator([rich]).eval_trend_new(rich, None)
        assert e1.admission_basis == "provisional_rule"
        assert (e1.admission_status, e1.grade, e1.score) \
            == (e2.admission_status, e2.grade, e2.score)


class TestBusinessConfirmationChannel:
    """规则 §7.3 A.6/B.3：业务人员确认是合法证据（configs/business_confirmations.yaml）。"""

    def _cfg_with(self, entries):
        c = copy.deepcopy(CFG)
        c["business_confirmations"] = {"confirmations": entries}
        return c

    def test_confirmed_novelty_and_freshness_upgrade_to_confirmed_basis(self):
        p = base_packet("bc1")
        cfg = self._cfg_with([
            {"candidate_id": "bc1", "field": "freshness", "status": "confirmed",
             "note": "业务确认近期推出"},
            {"candidate_id": "bc1", "field": "novelty", "status": "confirmed",
             "note": "业务确认新款型"}])
        evals, _, _ = TrackEvaluator(cfg, GroupStats([p]), "run_t").run([p])
        e = next(x for x in evals if x.track_id == "trend_new")
        assert e.admission_status == "ELIGIBLE"
        assert e.admission_basis == "confirmed"      # 不再是占位
        assert any(r.startswith("ev_bizconfirm_") for r in e.evidence_refs)

    def test_confirmed_old_style_is_rejected(self):
        p = base_packet("bc2")
        cfg = self._cfg_with([
            {"candidate_id": "bc2", "field": "novelty", "status": "old",
             "note": "业务确认旧款型"}])
        evals, _, _ = TrackEvaluator(cfg, GroupStats([p]), "run_t").run([p])
        e = next(x for x in evals if x.track_id == "trend_new")
        assert e.admission_status == "REJECTED"


class TestScoringDimsCoverage:
    """规则 §7.6/§8.5/§9.3：声明的评分维度必须全部出现在评分明细
    （已计分或留槽 0 分标注），保证每个分数可解释（§16.8）。"""

    def test_trend_components_cover_all_dims(self):
        p = base_packet("dc1")
        p.set_fact("basic_facts", "price", {"amount": 25.0, "currency": "USD"},
                   ev("basic_facts.price", "ev_dc1_p"))
        e = make_evaluator([p]).eval_trend_new(p, None)
        names = "".join(c["name"] for c in e.components)
        for kw in ("新鲜度", "共振", "内容扩散", "可定义", "市场适配"):
            assert kw in names, f"趋势评分明细缺维度：{kw}"
        fit = next(c for c in e.components if "市场适配" in c["name"])
        assert fit["score"] > 0        # 有价格即计分（价格带代理）

    def test_hit_components_cover_all_dims(self):
        p = base_packet("dc2", group="shein_frontend", roles=("transaction", "voc"))
        p.set_fact("market_metrics", "sales_floor_units", 800,
                   ev("market_metrics.sales_floor_units", "ev_dc2_s"))
        p.set_fact("basic_facts", "rating", 3.8, ev("basic_facts.rating", "ev_dc2_r"))
        p.set_fact("basic_facts", "review_count", 120,
                   ev("basic_facts.review_count", "ev_dc2_rc"))
        p.product_opportunity["negative_review_clusters"] = [
            {"cluster_name": f"痛点{i}", "mention_count": 5,
             "evidence_refs": [f"ev_dc2_c{i}"]} for i in range(3)]
        e = make_evaluator([p]).eval_hit_improvement(p)
        names = "".join(c["name"] for c in e.components)
        for kw in ("需求已验证", "痛点强度", "可改性", "差异化", "竞争与价格空间"):
            assert kw in names, f"改款评分明细缺维度：{kw}"

    def test_longtail_components_cover_all_dims_and_risk(self):
        p = base_packet("dc3", group="shein_frontend", roles=("transaction",))
        p.set_fact("market_metrics", "sales_floor_units", 400,
                   ev("market_metrics.sales_floor_units", "ev_dc3_s"))
        p.set_fact("basic_facts", "price", {"amount": 19.9, "currency": "USD"},
                   ev("basic_facts.price", "ev_dc3_p"))
        p.set_fact("competition_metrics", "seller_type", "BRAND",
                   ev("competition_metrics.seller_type", "ev_dc3_b"))
        e = make_evaluator([p]).eval_long_tail(p)
        names = "".join(c["name"] for c in e.components)
        for kw in ("需求稳定性", "竞争可突破", "价格带", "事实完整度", "直接承接"):
            assert kw in names, f"长尾评分明细缺维度：{kw}"
        risk = next(c for c in e.components if "风险扣减" in c["name"])
        assert risk["score"] < 0       # 品牌自营风险自动扣减（草案）

    def test_hit_relative_low_rating_uses_group_percentile(self):
        """低评分相对口径：组内 p25 之下即偏低，即使高于绝对线 4.0。"""
        peers = []
        for i in range(9):
            q = base_packet(f"pr{i}", group="shein_frontend",
                            roles=("transaction", "voc"))
            q.set_fact("basic_facts", "rating", 4.8,
                       ev("basic_facts.rating", f"ev_pr{i}_r"))
            peers.append(q)
        target = base_packet("prx", group="shein_frontend",
                             roles=("transaction", "voc"))
        target.set_fact("market_metrics", "sales_floor_units", 800,
                        ev("market_metrics.sales_floor_units", "ev_prx_s"))
        target.set_fact("basic_facts", "rating", 4.2,
                        ev("basic_facts.rating", "ev_prx_r"))
        target.set_fact("basic_facts", "review_count", 120,
                        ev("basic_facts.review_count", "ev_prx_rc"))
        evaluator = make_evaluator(peers + [target])
        evaluator._peers["shein_frontend"] = peers + [target]
        e = evaluator.eval_hit_improvement(target)
        # 4.2 高于绝对线 4.0，但低于组内 p25(4.8) -> 相对偏低成立，可占位准入
        assert e.admission_status == "ELIGIBLE"
        assert e.admission_basis == "provisional_rule"


class TestTrackReportGovernance:
    """《执行矛盾与流程待确认问题》临时执行规则在交付表上的落地。"""

    @pytest.fixture()
    def workbook(self, tmp_path):
        from openpyxl import load_workbook
        from grading_system.xlsx_report import export_track_report
        packets = [base_packet(f"x{i}") for i in range(3)]
        packets[0].freshness.update({"freshness_status": "confirmed",
                                     "new_arrival_flag": "page_new_arrival",
                                     "novelty_status": "confirmed",
                                     "freshness_evidence_refs": ["ev_f"],
                                     "novelty_evidence_refs": ["ev_n"]})
        evals, clusters, budgets = make_evaluator(packets).run(packets)
        out = tmp_path / "三赛道.xlsx"
        export_track_report(out, packets, evals, clusters, budgets, CFG)
        wb = load_workbook(out, read_only=True)
        yield wb
        wb.close()

    def test_p0_04_no_cross_track_auto_decision(self, workbook):
        """跨赛道汇总只并列展示各赛道开发方式，不自动选「最佳赛道」。"""
        headers = [str(c.value or "")
                   for c in next(workbook["跨赛道汇总"].iter_rows(max_row=1))]
        assert any("人工决策" in h for h in headers)
        assert not any(("建议赛道" in h) or ("最佳" in h) for h in headers)

    def test_p0_03_pending_sheet_marks_refetch_budget(self, workbook):
        """待补数据 sheet 标注每行是否进入本轮补证预算队列。"""
        rows = list(workbook["待补数据"].iter_rows(values_only=True))
        assert "本轮补证预算" in rows[0]
        col = rows[0].index("本轮补证预算")
        marks = {r[col] for r in rows[1:] if r[col]}
        assert marks <= {"入队(open)", "待下轮"} and "入队(open)" in marks

    def test_p0_09_rule_sheet_carries_project_status_and_refetch(self, workbook):
        keys = [str(r[0]) for r in
                workbook["规则版本与运行记录"].iter_rows(values_only=True)]
        assert "project_status" in keys
        assert any("refetch_queues" in k for k in keys)
