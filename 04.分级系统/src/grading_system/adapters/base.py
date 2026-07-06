"""Adapter 基类：XLSX 读取 + SourceEnvelope / EvidenceRef 工厂。"""
from __future__ import annotations

import itertools
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, Iterator, List, Optional

from openpyxl import load_workbook

from ..models import EvidenceRef, SourceEnvelope

COLLECTOR_VERSION = "grading-system/0.1.0"

_evidence_counter = itertools.count(1)


def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def read_sheet_dicts(workbook_path: Path, sheet: str) -> List[Dict[str, Any]]:
    """把一个 sheet 读成 dict 列表，附带 ``_row``（1-based 数据行号，含表头偏移）。"""
    wb = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        ws = wb[sheet]
        rows = ws.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            return []
        header = [str(h).strip() if h is not None else f"col_{i}" for i, h in enumerate(header)]
        out = []
        for idx, row in enumerate(rows, start=2):
            if row is None or all(v is None for v in row):
                continue
            d = {header[i]: row[i] for i in range(min(len(header), len(row)))}
            d["_row"] = idx
            out.append(d)
        return out
    finally:
        wb.close()


class BaseAdapter:
    """所有数据源适配器的公共部分。"""

    source_id: str = "base"
    source_type: str = "plugin"

    def __init__(self, workbook_path: Path, market: str = "US",
                 collection_mode: str = "full_then_analyze", layer_hint: str = "L1"):
        self.workbook_path = Path(workbook_path)
        self.market = market
        self.collection_mode = collection_mode
        self.layer_hint = layer_hint
        self.warnings: List[str] = []
        self.errors: List[str] = []

    # ---------- 工厂 ----------

    def make_evidence(self, field_path: str, sheet: str, row: Any,
                      confidence: str = "high", note: str = "",
                      observed_at: Optional[str] = None,
                      source_type: Optional[str] = None) -> EvidenceRef:
        return EvidenceRef(
            evidence_id=f"ev_{self.source_id}_{next(_evidence_counter):06d}",
            source_id=self.source_id,
            source_type=source_type or self.source_type,
            tool_or_adapter=type(self).__name__,
            field_path=field_path,
            artifact_path=str(self.workbook_path),
            record_locator=f"sheet={sheet};row={row}",
            observed_at=observed_at,
            confidence=confidence,
            note=note,
        )

    def make_envelope(self, records: List[Dict[str, Any]],
                      quality_status: str = "VALID") -> SourceEnvelope:
        return SourceEnvelope(
            envelope_id=f"envelope_{self.source_id}_{datetime.now(timezone.utc):%Y%m%d}",
            source_id=self.source_id,
            source_type=self.source_type,
            market=self.market,
            collection_mode=self.collection_mode,
            layer_hint=self.layer_hint,
            collected_at=now_iso(),
            collector_version=COLLECTOR_VERSION,
            quality_status=quality_status,
            records=records,
            errors=list(self.errors),
            warnings=list(self.warnings),
            artifact_paths={"raw_workbook": str(self.workbook_path)},
        )

    def collect(self) -> SourceEnvelope:  # pragma: no cover - 由子类实现
        raise NotImplementedError
