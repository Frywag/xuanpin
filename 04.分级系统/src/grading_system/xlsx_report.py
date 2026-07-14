"""XLSX 推荐表导出（只读 AnalysisResult / CandidateDataPacket，不触发新采集）。

Sheet 结构对应《07_分层分析系统组任务书》§12：
  推荐总表 / 评分明细 / 数据源覆盖 / 差评改良 / 多平台比对 / 风险清单 /
  证据索引 / 运行记录
"""
from __future__ import annotations

import json
from typing import Any, Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import AnalysisResult, CandidateDataPacket
from .packet_builder import KEY_FIELDS

GRADE_COLORS = {"S": "FFB45309", "A": "FF2563EB", "B": "FF6B7280", "C": "FFD1D5DB"}
HEADER_FILL = PatternFill("solid", fgColor="FF1F2937")
HEADER_FONT = Font(color="FFFFFFFF", bold=True)


def _money(v) -> str:
    if not v or not isinstance(v, dict) or v.get("amount") is None:
        return ""
    cur = v.get("currency") or "?"
    return f"{v['amount']:,.2f} {cur}"


def _sheet(wb: Workbook, title: str, headers: List[str], widths: List[int]):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    return ws


def export_track_report(path, packets, evaluations, clusters, budgets, tracks_cfg):
    """三赛道独立推荐工作簿（tracks.v1，校准态）。

    页签对应《三赛道独立SAB分级》§13：三张 SAB 表、跨赛道汇总、待补数据、
    未准入与淘汰依据、机会簇、数据源覆盖与角色、规则版本与运行记录。
    评分明细/证据索引/152键在主工作簿（选品推荐表.xlsx），此处不重复。
    """
    wb = Workbook()
    wb.remove(wb.active)
    by_id = {p.candidate_id: p for p in packets}
    labels = {tid: t["label"] for tid, t in tracks_cfg["tracks"].items()}

    def cand_cols(e):
        p = by_id.get(e.subject_id)
        return [e.subject_id, p.basic_facts.get("title") if p else "",
                p.platform if p else "",
                ",".join(p.context.get("source_roles") or []) if p else ""]

    for tid, label in labels.items():
        ws = _sheet(wb, f"{label}_SAB",
                    ["rank", "grade", "candidate_id", "title", "platform",
                     "source_roles", "score", "准入理由", "confidence", "风险",
                     "缺失字段", "下一步补采", "cluster_id", "rule_status"],
                    [6, 7, 30, 46, 14, 20, 8, 60, 10, 30, 36, 40, 16, 26])
        eligible = sorted([e for e in evaluations
                           if e.track_id == tid and e.admission_status == "ELIGIBLE"],
                          key=lambda e: -(e.score or 0))
        for i, e in enumerate(eligible, 1):
            ws.append([i, e.grade] + cand_cols(e) + [
                e.score, "；".join(e.admission_reasons)[:300], e.confidence,
                "；".join(e.risks)[:120], "、".join(e.missing_fields)[:150],
                "；".join(e.refetch_tasks[:3]), e.cluster_id or "", e.rule_status])
        ws.auto_filter.ref = ws.dimensions

    ws = _sheet(wb, "跨赛道汇总",
                ["candidate_id", "title", "platform",
                 "趋势新品", "爆款改款", "长尾直接选品", "可选开发方式(并列展示,人工决策)"],
                [30, 46, 14, 22, 22, 22, 40])
    by_cand = {}
    for e in evaluations:
        by_cand.setdefault(e.subject_id, {})[e.track_id] = e
    for cid, m in by_cand.items():
        if not any(e.admission_status == "ELIGIBLE" for e in m.values()):
            continue
        p = by_id.get(cid)
        def cell(tid):
            e = m.get(tid)
            if e is None:
                return ""
            return e.grade or e.admission_status
        # P0-04 临时口径：三赛道分数未经金标统一标尺，禁止跨赛道直接比较、
        # 禁止用最高分自动带出开发方式——并列展示全部已准入赛道的开发方式，人工决策
        modes = "；".join(f"{labels[e.track_id]}→{e.recommended_development_mode}"
                          for e in m.values() if e.grade)
        ws.append([cid, p.basic_facts.get("title") if p else "",
                   p.platform if p else "", cell("trend_new"),
                   cell("hit_improvement"), cell("long_tail_direct"), modes])
    ws.auto_filter.ref = ws.dimensions

    # 补证预算（P0-03）：进入 refetch_queue 的候选获得本轮补采名额，
    # 其余 PENDING 留在待补池等下轮预算——避免「缺数据者永远无预算」死锁
    in_refetch = {(tid, item["subject_id"])
                  for tid, b in budgets.items()
                  for item in b.get("refetch_queue", [])}
    ws = _sheet(wb, "待补数据",
                ["track", "candidate_id", "title", "缺失证据", "补采任务",
                 "本轮补证预算"],
                [16, 30, 44, 60, 50, 14])
    for e in evaluations:
        if e.admission_status == "PENDING_DATA":
            p = by_id.get(e.subject_id)
            ws.append([labels[e.track_id], e.subject_id,
                       p.basic_facts.get("title") if p else "",
                       "、".join(e.missing_fields)[:200],
                       "；".join(e.refetch_tasks[:3]),
                       "入队(open)" if (e.track_id, e.subject_id) in in_refetch
                       else "待下轮"])
    ws.auto_filter.ref = ws.dimensions

    ws = _sheet(wb, "未准入与淘汰依据",
                ["track", "candidate_id", "title", "淘汰依据", "evidence_refs"],
                [16, 30, 44, 70, 36])
    for e in evaluations:
        if e.admission_status == "REJECTED":
            p = by_id.get(e.subject_id)
            ws.append([labels[e.track_id], e.subject_id,
                       p.basic_facts.get("title") if p else "",
                       "；".join(e.admission_reasons)[:250],
                       ",".join(e.evidence_refs[:5])])
    ws.auto_filter.ref = ws.dimensions

    ws = _sheet(wb, "机会簇与跨平台相似",
                ["cluster_id", "version", "代表款型", "关系类型", "confidence",
                 "human_review", "成员数", "成员（candidate_id×source_group）"],
                [16, 8, 30, 12, 11, 13, 8, 80])
    for c in clusters:
        ws.append([c["cluster_id"], c["cluster_version"], c["representative_style"],
                   c["relation_type"], c["confidence"], c["human_review"],
                   len(c["members"]),
                   "；".join(f"{m['candidate_id']}({m['source_group']})"
                             for m in c["members"][:8])])
    ws.auto_filter.ref = ws.dimensions

    ws = _sheet(wb, "数据源覆盖与角色",
                ["source_group", "候选数", "roles", "趋势准入判定覆盖"],
                [28, 10, 30, 22])
    from collections import Counter as _C
    grp = _C(p.context.get("source_group") for p in packets)
    roles_of = {p.context.get("source_group"):
                ",".join(p.context.get("source_roles") or []) for p in packets}
    for g, n in grp.most_common():
        ws.append([g, n, roles_of.get(g, ""), "100%（每候选三条赛道记录）"])

    ws = _sheet(wb, "规则版本与运行记录", ["key", "value"], [36, 100])
    for k, v in [("rule_version", tracks_cfg.get("rule_version")),
                 ("rule_status", tracks_cfg.get("rule_status")),
                 ("project_status", tracks_cfg.get("project_status")),
                 ("说明", "所有等级为校准态草案，业务金标与参数批准前不是正式生产等级"),
                 ("evaluations", len(evaluations)),
                 ("clusters", len(clusters)),
                 ("l2_queues", json.dumps({k: len(v["l2_queue"]) for k, v in budgets.items()},
                                          ensure_ascii=False)),
                 ("l3_queues", json.dumps({k: len(v["l3_queue"]) for k, v in budgets.items()},
                                          ensure_ascii=False)),
                 ("refetch_queues(补证预算,P0-03)",
                  json.dumps({k: len(v.get("refetch_queue", []))
                              for k, v in budgets.items()}, ensure_ascii=False)),
                 ("待业务确认参数", "；".join(tracks_cfg.get(
                     "pending_business_confirmation", [])))]:
        ws.append([k, str(v)])
    wb.save(path)
    return path


def export_report(path, packets: List[CandidateDataPacket],
                  results: Dict[str, AnalysisResult],
                  run_meta: Dict[str, Any],
                  evidence_sheet_top_n: int = 300):
    wb = Workbook()
    wb.remove(wb.active)
    by_id = {p.candidate_id: p for p in packets}
    ordered = sorted(results.values(),
                     key=lambda r: -(r.priority_score["pct"] or 0))

    # ---------- 1. 推荐总表 ----------
    ws = _sheet(wb, "推荐总表",
                ["candidate_id", "product_link", "image", "title", "category",
                 "source_platform", "track", "grade", "priority_score", "pct",
                 "track_tags", "recommendation_reason",
                 "risk_summary", "missing_key_fields", "next_action", "evidence_count",
                 "confidence", "L1", "L2", "L3"],
                [30, 40, 40, 50, 20, 16, 12, 7, 12, 8, 22, 60, 50, 40, 40, 12, 10, 6, 6, 6])
    for r in ordered:
        p = by_id[r.candidate_id]
        category = (p.context.get("category_name")
                    or "/".join(p.basic_facts.get("category_path") or [])
                    or p.context.get("category_label") or "")
        ws.append([
            r.candidate_id, p.canonical_url or "", p.basic_facts.get("image_url") or "",
            p.basic_facts.get("title") or "", category, p.platform,
            r.priority_score.get("track_label") or "",
            r.priority_score["grade"] or "", r.priority_score["total"],
            r.priority_score["pct"], "、".join(r.track_tags),
            r.recommendation["reason_summary"], r.recommendation["risk_summary"],
            "、".join(r.missing_fields[:6]), "；".join(r.recommendation["next_actions"][:3]),
            len(p.evidence_pack), r.confidence,
            r.layer_results["L1"].grade or "", r.layer_results["L2"].grade or "",
            r.layer_results["L3"].grade or ""])
        g = r.priority_score["grade"]
        if g in GRADE_COLORS:
            ws.cell(row=ws.max_row, column=8).font = Font(
                color=GRADE_COLORS[g], bold=g in ("S", "A"))
    ws.auto_filter.ref = ws.dimensions

    # ---------- 2. 评分明细 ----------
    ws = _sheet(wb, "评分明细",
                ["candidate_id", "grade", "维度", "得分", "满分", "理由",
                 "evidence_refs", "missing"],
                [30, 7, 12, 8, 8, 70, 45, 8])
    for r in ordered:
        for c in r.priority_score["components"]:
            ws.append([r.candidate_id, r.priority_score["grade"], c.name,
                       c.score, c.max_score, c.reason,
                       ",".join(c.evidence_refs[:8]), "Y" if c.missing else ""])
    ws.auto_filter.ref = ws.dimensions

    # ---------- 3. 数据源覆盖 ----------
    ws = _sheet(wb, "数据源覆盖",
                ["candidate_id", "sources", "source_count", "关键字段覆盖率",
                 "quality_status", "missing_fields", "field_issues"],
                [30, 35, 12, 15, 18, 60, 60])
    for r in ordered:
        p = by_id[r.candidate_id]
        ws.append([p.candidate_id,
                   "、".join(sorted({s['source_id'] for s in p.source_refs})),
                   p.source_count, round(p.field_coverage(KEY_FIELDS), 2),
                   p.quality_status, "、".join(p.missing_fields),
                   "；".join(f"{i.field_path}:{i.issue}" for i in p.field_issues[:5])])
    ws.auto_filter.ref = ws.dimensions

    # ---------- 4. 差评改良 ----------
    ws = _sheet(wb, "差评改良",
                ["candidate_id", "聚类", "提及次数", "频率", "代表标签",
                 "改良方向", "不改的风险", "evidence_refs"],
                [30, 16, 10, 8, 40, 30, 30, 40])
    for r in ordered:
        p = by_id[r.candidate_id]
        for c in p.product_opportunity.get("negative_review_clusters", []):
            ws.append([p.candidate_id, c["cluster_name"], c["mention_count"],
                       c["frequency_hint"], "、".join(c["representative_reviews"]),
                       c["improvement_opportunity"], c["risk_if_unfixed"],
                       ",".join(c["evidence_refs"][:5])])

    # ---------- 4b. 商品画像（152 键模板；直取+manual，语义键经 LLM 回灌入库） ----------
    ws = _sheet(wb, "商品画像",
                ["candidate_id", "分组", "键", "值", "来源", "evidence_refs"],
                [30, 14, 18, 60, 12, 40])
    for r in ordered:
        p = by_id[r.candidate_id]
        for name, item in (p.context.get("profile") or {}).items():
            ws.append([p.candidate_id, item.get("group", ""), name,
                       str(item.get("value")) if item.get("value") is not None
                       else "（待人工补录）",
                       item.get("source", ""),
                       ",".join(item.get("evidence_refs", [])[:4])])
    ws.auto_filter.ref = ws.dimensions

    # ---------- 5. 多平台比对 ----------
    ws = _sheet(wb, "多平台比对",
                ["candidate_id", "类型", "内容", "币种", "备注"],
                [30, 22, 70, 10, 50])
    for r in ordered:
        p = by_id[r.candidate_id]
        for m in p.cross_platform.get("matched_products", []):
            ws.append([p.candidate_id, "多源商品ID互验",
                       f"{m['source_id']} / product_id={m['product_id']}", "",
                       m["match_type"]])
        for row in p.cross_platform.get("platform_price_comparison", []):
            ws.append([p.candidate_id, f"价格带（{row['style_group']}）",
                       f"{row['source_group']}: p25={row['p25']} p50={row['p50']} "
                       f"p75={row['p75']} (n={row['n']})",
                       row["currency"] or "", row["note"]])
        for s in p.cross_platform.get("same_style_signals", []):
            ws.append([p.candidate_id, "同风格信号", s["signal"], "", s["note"]])

    # ---------- 6. 风险清单 ----------
    ws = _sheet(wb, "风险清单",
                ["candidate_id", "grade", "风险描述", "confidence", "evidence_refs"],
                [30, 7, 80, 12, 45])
    for r in ordered:
        for c in r.claims:
            if c.claim_type == "risk":
                ws.append([r.candidate_id, r.priority_score["grade"], c.claim,
                           c.confidence, ",".join(c.evidence_refs[:6])])
    ws.auto_filter.ref = ws.dimensions

    # ---------- 7. 证据索引（前 N 个候选，全量证据见 JSON 产物） ----------
    ws = _sheet(wb, "证据索引",
                ["evidence_id", "candidate_id", "source_id", "field_path",
                 "artifact_path", "record_locator", "confidence", "note"],
                [26, 30, 16, 40, 55, 28, 10, 40])
    for r in ordered[:evidence_sheet_top_n]:
        p = by_id[r.candidate_id]
        for e in p.evidence_pack:
            ws.append([e.evidence_id, p.candidate_id, e.source_id, e.field_path,
                       e.artifact_path, e.record_locator, e.confidence, e.note])
    ws.auto_filter.ref = ws.dimensions

    # ---------- 8. 运行记录 ----------
    ws = _sheet(wb, "运行记录", ["key", "value"], [40, 110])
    ws.append(["结果口径", "本工作簿为 Legacy 基线结果（旧单赛道链路），非正式生产推荐；"
               "三赛道校准结果见 三赛道选品推荐表.xlsx 与 track_evaluations 表"])
    for k, v in run_meta.items():
        ws.append([k, str(v)])

    wb.save(path)
    return path
